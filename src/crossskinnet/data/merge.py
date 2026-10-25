"""Schema-driven raw dataset ingestion with strictly separate pool outputs."""
from __future__ import annotations

import csv
from collections import Counter
from pathlib import Path
from typing import Any

import yaml

BASE_COLUMNS = ("image_id", "image_path", "source_dataset", "original_label",
                "patient_id", "lesion_id", "age", "sex", "site")
TRAIN = ("PAD-UFES-20", "ISIC2018", "Dermnet")
EXTERNAL = ("HAM10000", "PH2")
MISSING = {"", "nan", "na", "null", "none"}


def read_yaml(path: Path) -> dict[str, Any]:
    """Read a YAML mapping without inferring missing configuration."""
    with path.open() as handle:
        result = yaml.safe_load(handle)
    if not isinstance(result, dict):
        raise ValueError(f"Expected a YAML mapping in {path}")
    return result


def _load_csv(root: Path, schema: dict[str, Any], name: str) -> list[dict[str, str]]:
    """Load explicitly listed metadata files and column mappings, without split filtering."""
    if not schema:
        raise ValueError(f"{name}: schema unavailable. Supply real inspection output")
    if schema.get("format") == "folders":
        return _folder_rows(root, schema, name)
    if schema.get("format") == "xlsx":
        return _spreadsheet_rows(root, schema, name)
    if schema.get("format") != "csv":
        raise ValueError(f"{name}: unsupported configured format")
    files = schema.get("metadata_files")
    columns = schema.get("columns", {})
    if not isinstance(files, list) or not files or not isinstance(columns, dict):
        raise ValueError(f"{name}: require metadata_files list and columns mapping")
    for field in ("image_id", "image_path", "original_label"):
        if field == "original_label" and schema.get("one_hot_labels"):
            continue
        if not columns.get(field):
            raise ValueError(f"{name}: missing required column mapping {field}")
    symptoms = schema.get("symptoms", {})
    mapping = {**columns, **symptoms}
    allowed = set(BASE_COLUMNS) - {"source_dataset"}
    if set(columns) - allowed or set(symptoms) & set(BASE_COLUMNS):
        raise ValueError(f"{name}: invalid or conflicting output column mappings")
    rows = []
    image_index = _image_index(root, schema) if schema.get("image_globs") else {}
    join = schema.get("join")
    joined = {}
    if join:
        with (root / join["file"]).open(newline="", encoding="utf-8-sig") as handle:
            reader = csv.DictReader(handle)
            required = {join["key"], *join["columns"].values()}
            if not required <= set(reader.fieldnames or []):
                raise ValueError(f"{name}: missing configured join columns")
            for record in reader:
                key = record[join["key"]]
                if key in joined:
                    raise ValueError(f"{name}: duplicate join key {key}")
                joined[key] = record
    for filename in files:
        metadata = root / filename
        with metadata.open(newline="", encoding=schema.get("encoding", "utf-8-sig")) as handle:
            reader = csv.DictReader(handle, delimiter=schema.get("delimiter", ","))
            missing = (set(mapping.values()) | set(schema.get("one_hot_labels", []))) - set(reader.fieldnames or [])
            if missing:
                raise ValueError(f"{name}: {metadata} missing configured columns: {sorted(missing)}")
            for line, raw in enumerate(reader, 2):
                row = {field: "NaN" for field in (*BASE_COLUMNS, *symptoms)}
                row["source_dataset"] = name
                for field, source in mapping.items():
                    value = (raw[source] or "").strip()
                    row[field] = "NaN" if value.lower() in MISSING | {v.lower() for v in schema.get("missing_values", [])} else value
                if schema.get("one_hot_labels"):
                    labels = schema["one_hot_labels"]
                    values = [float(raw[label]) for label in labels]
                    if any(value not in (0, 1) for value in values) or sum(values) != 1:
                        raise ValueError(f"{name}: expected exactly one active label at {metadata}:{line}")
                    row["original_label"] = labels[values.index(1)]
                if join:
                    key = raw[join["key"]]
                    if key not in joined:
                        raise ValueError(f"{name}: missing lesion grouping for {key}")
                    for field, column in join["columns"].items():
                        value = joined[key][column].strip()
                        row[field] = "NaN" if value.lower() in MISSING else value
                for field in ("image_id", "image_path", "original_label"):
                    if row[field] == "NaN":
                        raise ValueError(f"{name}: missing {field} in {metadata}:{line}")
                value = row["image_path"]
                if schema.get("image_globs"):
                    if value not in image_index:
                        raise ValueError(f"{name}: no configured image for {value}")
                    image = image_index[value]
                else:
                    image = (root / schema.get("image_template", "{value}").format(value=value)).resolve()
                if not image.is_file():
                    raise ValueError(f"{name}: configured image does not exist: {image}")
                row["image_path"] = str(image)
                rows.append(row)
    if not rows:
        raise ValueError(f"{name}: configured metadata contains no images")
    return rows


def _configured_images(root: Path, schema: dict[str, Any]) -> list[Path]:
    """Enumerate only explicitly configured image locations, deterministically."""
    paths = sorted({path.resolve() for pattern in schema["image_globs"]
                    for path in root.glob(pattern) if path.is_file()})
    extensions = schema.get("extensions")
    if extensions:
        paths = [path for path in paths if path.suffix.lower() in extensions]
    if not paths:
        raise ValueError(f"No images matched configured globs under {root}")
    return paths


def _image_index(root: Path, schema: dict[str, Any]) -> dict[str, Path]:
    """Index IDs by filename stem, rejecting ambiguous physical copies."""
    index: dict[str, Path] = {}
    for path in _configured_images(root, schema):
        if path.stem in index:
            raise ValueError(f"Ambiguous image ID {path.stem}: {index[path.stem]} and {path}")
        index[path.stem] = path
    return index


def _folder_rows(root: Path, schema: dict[str, Any], name: str) -> list[dict[str, str]]:
    """Preserve class folder names across every configured original split."""
    if schema.get("label_from") != "parent_name":
        raise ValueError(f"{name}: require explicit label_from: parent_name")
    rows = []
    for image in _configured_images(root, schema):
        row = dict.fromkeys(BASE_COLUMNS, "NaN")
        row.update(image_id=image.relative_to(root.resolve()).as_posix(),
                   image_path=str(image), source_dataset=name,
                   original_label=image.parent.name)
        rows.append(row)
    return rows


def _spreadsheet_rows(root: Path, schema: dict[str, Any], name: str) -> list[dict[str, str]]:
    """Find inspected PH2 headers and clinical X markers, excluding masks/ROIs."""
    from openpyxl import load_workbook

    images = _image_index(root, schema)
    workbook = load_workbook(root / schema["metadata_file"], read_only=True, data_only=True)
    try:
        records = list(workbook[schema["sheet"]].iter_rows(values_only=True))
    finally:
        workbook.close()
    headers = [schema["image_header"], *schema["label_headers"]]
    positions: dict[str, int] = {}
    for header in headers:
        candidates = records if header == schema["image_header"] else [
            record for record in records
            if str(record[positions[schema["image_header"]]] or "").strip() not in images]
        matches = {index for record in candidates for index, value in enumerate(record)
                   if isinstance(value, str) and value.strip() == header}
        if len(matches) != 1:
            raise ValueError(f"{name}: need an unambiguous spreadsheet column for {header!r}")
        positions[header] = matches.pop()
    rows = []
    seen = set()
    for record in records:
        value = str(record[positions[schema["image_header"]]] or "").strip()
        if value not in images:
            continue
        if value in seen:
            raise ValueError(f"{name}: duplicate spreadsheet image ID {value}")
        seen.add(value)
        labels = [label for label in schema["label_headers"]
                  if str(record[positions[label]] or "").strip() == schema["label_marker"]]
        if len(labels) != 1:
            raise ValueError(f"{name}: expected exactly one clinical label for {value}")
        row = dict.fromkeys(BASE_COLUMNS, "NaN")
        row.update(image_id=value, image_path=str(images[value]),
                   source_dataset=name, original_label=labels[0])
        rows.append(row)
    if set(images) != seen:
        raise ValueError(f"{name}: images missing spreadsheet labels: {sorted(set(images) - seen)[:10]}")
    return rows


def load_pad_ufes20(root: Path, schema: dict[str, Any]) -> list[dict[str, str]]:
    """Load PAD-UFES-20 using inspected schema configuration."""
    return _load_csv(root, schema, "PAD-UFES-20")


def load_isic2018(root: Path, schema: dict[str, Any]) -> list[dict[str, str]]:
    """Load every configured ISIC2018 metadata partition."""
    return _load_csv(root, schema, "ISIC2018")


def load_dermnet(root: Path, schema: dict[str, Any]) -> list[dict[str, str]]:
    """Load Dermnet using inspected schema configuration."""
    return _load_csv(root, schema, "Dermnet")


def load_ham10000(root: Path, schema: dict[str, Any]) -> list[dict[str, str]]:
    """Load HAM10000 exclusively for the external pool."""
    return _load_csv(root, schema, "HAM10000")


def load_ph2(root: Path, schema: dict[str, Any]) -> list[dict[str, str]]:
    """Load PH2 exclusively for the external pool."""
    return _load_csv(root, schema, "PH2")


LOADERS = dict(zip((*TRAIN, *EXTERNAL),
                   (load_pad_ufes20, load_isic2018, load_dermnet, load_ham10000, load_ph2)))


def merge_datasets(config_path: Path, schemas_path: Path) -> tuple[Path, Path]:
    """Validate all datasets, then write raw pools without deduplication or splitting."""
    config = read_yaml(config_path)
    schemas = read_yaml(schemas_path)["datasets"]
    pools: dict[str, list[dict[str, str]]] = {"train": [], "external_test": []}
    for name, loader in LOADERS.items():
        settings = config["datasets"][name]
        expected = "train" if name in TRAIN else "external_test"
        if settings.get("pool") != expected:
            raise ValueError(f"{name}: pool must be {expected}; external data cannot train/validate")
        root = Path(settings["path"])
        if not root.is_absolute():
            root = config_path.resolve().parent / root
        pools[expected].extend(loader(root, schemas.get(name)))
    symptoms = sorted({key for rows in pools.values() for row in rows for key in row} - set(BASE_COLUMNS))
    fields = [*BASE_COLUMNS, *symptoms]
    metadata_fields = ["patient_id", "lesion_id", "age", "sex", "site", *symptoms]
    fields.extend(f"{field}_mask" for field in metadata_fields)
    output = Path(config["paths"]["output_root"])
    if not output.is_absolute():
        output = config_path.resolve().parent / output
    output.mkdir(parents=True, exist_ok=True)
    paths = (output / "raw_train_pool.csv", output / "raw_test_pool.csv")
    for (pool, rows), path in zip(pools.items(), paths):
        for row in rows:
            for field in metadata_fields:
                row.setdefault(field, "NaN")
                row[f"{field}_mask"] = "0" if row[field] == "NaN" else "1"
        with path.open("w", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=fields)
            writer.writeheader()
            writer.writerows(rows)
        print(f"{pool}: shape=({len(rows)}, {len(fields)}); head={rows[:2]}")
        print(f"Counts per source_dataset: {dict(Counter(row['source_dataset'] for row in rows))}")
        print(f"Counts per original_label: {dict(Counter(row['original_label'] for row in rows))}")
        print(f"Saved {path}")
    return paths
